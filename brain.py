"""Jarvis brain: models, tools, memory. No UI code here."""
import base64
import datetime
import io
import json
import os
import re
import urllib.parse
import webbrowser
from pathlib import Path

from dotenv import load_dotenv
from openai import OpenAI

HERE = Path(__file__).parent
load_dotenv(HERE / ".env")

MEMORY_FILE = HERE / "memory.json"
DOCS_DIR = Path.home() / "Documents" / "Jarvis"
HISTORY_LIMIT = 60          # messages (incl. tool calls/results) sent to the model
HISTORY_TOOL_CHARS = 400    # tool results are cut to this in history so file reads don't flood it
URL_LIMIT = 30              # remembered sites shown to the model
MAX_TOOL_ROUNDS = 5
MAX_SCREEN_STEPS = 15
SCREENSHOT_WIDTH = 1280

# File tools only touch these folders: your user folder plus JARVIS_FILE_ROOTS (separated by ';').
FILE_ROOTS = [Path.home()] + [Path(p) for p in os.getenv("JARVIS_FILE_ROOTS", "").split(os.pathsep) if p.strip()]
BACKUP_DIR = DOCS_DIR / "backups"
MAX_READ_CHARS = 20000      # keeps a big file from flooding the small model's context
MAX_LIST_ENTRIES = 100

local = OpenAI(base_url=os.getenv("OLLAMA_URL", "http://localhost:11434/v1"), api_key="ollama")
# Non-thinking 4B: fits fully in a 6 GB GPU and answers in ~1s; thinking models took 10-150s.
LOCAL_MODEL = os.getenv("LOCAL_MODEL", "qwen3:4b-instruct")

AZURE_ENDPOINT = os.getenv("AZURE_ENDPOINT", "").rstrip("/")
AZURE_MODEL = os.getenv("AZURE_MODEL", "")
azure = (OpenAI(base_url=AZURE_ENDPOINT + "/openai/v1/", api_key=os.getenv("AZURE_API_KEY", ""))
         if AZURE_ENDPOINT and AZURE_MODEL else None)

APP_ALIASES = {
    "word": "winword", "excel": "excel", "powerpoint": "powerpnt", "outlook": "outlook",
    "notepad": "notepad", "calculator": "calc", "hesap makinesi": "calc",
    "chrome": "chrome", "edge": "msedge", "explorer": "explorer", "paint": "mspaint",
}


# ---------- memory ----------

def load_memory(path=MEMORY_FILE):
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (FileNotFoundError, json.JSONDecodeError):
        return {"history": [], "urls": [], "facts": []}


def recent(history, limit=HISTORY_LIMIT):
    """Last `limit` messages, starting at a user turn so no tool result loses its tool call."""
    tail = history[-limit:]
    while tail and tail[0]["role"] != "user":
        tail = tail[1:]
    return tail


def save_memory(mem, path=MEMORY_FILE):
    trimmed = {**mem, "history": recent(mem["history"]), "urls": mem["urls"][-URL_LIMIT:]}
    path.write_text(json.dumps(trimmed, ensure_ascii=False, indent=1), encoding="utf-8")


mem = load_memory()


# ---------- tools ----------

def open_app(name):
    os.startfile(APP_ALIASES.get(name.lower().strip(), name))
    return f"Opened {name}."


def open_url(url):
    if not url.startswith(("http://", "https://")):
        url = "https://" + url
    if urllib.parse.urlparse(url).scheme not in ("http", "https"):
        return "Refused: only http/https links."
    webbrowser.open(url)
    mem["urls"].append({"url": url, "when": datetime.datetime.now().isoformat(timespec="minutes")})
    save_memory(mem)
    return f"Opened {url}."


def youtube_search(query):
    return open_url("https://www.youtube.com/results?search_query=" + urllib.parse.quote_plus(query))


def web_search(query):
    return open_url("https://www.google.com/search?q=" + urllib.parse.quote_plus(query))


def safe_filename(title):
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", title).strip(" .")
    return (name or "Jarvis document")[:80]


def write_word_doc(title, content):
    import docx  # python-docx; imported lazily so tests don't need it
    doc = docx.Document()
    doc.add_heading(title, level=0)
    for line in content.splitlines():
        if line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith(("- ", "* ")):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line.strip():
            doc.add_paragraph(line)
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    path = DOCS_DIR / f"{safe_filename(title)}.docx"
    doc.save(path)
    os.startfile(path)
    return f"Saved and opened {path}."


def draft_email(to, subject, body):
    # Only opens a draft in the default mail app; the user presses Send.
    query = urllib.parse.urlencode({"subject": subject, "body": body}, quote_via=urllib.parse.quote)
    webbrowser.open(f"mailto:{urllib.parse.quote(to)}?{query}")
    return "Email draft opened. The user must review and press Send."


def user_path(path):
    """Relative paths start at the user folder. Refuses anything outside FILE_ROOTS."""
    p = (Path.home() / os.path.expandvars(path.strip().strip('"'))).expanduser().resolve()
    if not any(p.is_relative_to(root.resolve()) for root in FILE_ROOTS):
        raise PermissionError(f"{p} is outside the allowed folders: {', '.join(map(str, FILE_ROOTS))}")
    return p


def backup(p):
    """Copy a file before Jarvis changes it, so nothing is lost."""
    import shutil
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.datetime.now().strftime("%Y%m%d-%H%M%S-%f")
    shutil.copy2(p, BACKUP_DIR / f"{stamp}_{p.name}")


def list_files(folder):
    p = user_path(folder)
    entries = sorted(p.iterdir(), key=lambda e: (e.is_file(), e.name.lower()))
    lines = [e.name + ("\\" if e.is_dir() else "") for e in entries[:MAX_LIST_ENTRIES]]
    more = f"\n…and {len(entries) - MAX_LIST_ENTRIES} more" if len(entries) > MAX_LIST_ENTRIES else ""
    return f"{p}:\n" + ("\n".join(lines) or "(empty)") + more


def read_file(path):
    p = user_path(path)
    text = p.read_text(encoding="utf-8", errors="replace")
    cut = f"\n…[cut, {len(text) - MAX_READ_CHARS} more chars]" if len(text) > MAX_READ_CHARS else ""
    return text[:MAX_READ_CHARS] + cut


def write_file(path, content):
    p = user_path(path)
    existed = p.exists()
    if existed:
        backup(p)
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(content, encoding="utf-8")
    return f"{'Overwrote' if existed else 'Created'} {p}."


def edit_file(path, find, replace):
    p = user_path(path)
    text = p.read_text(encoding="utf-8")
    count = text.count(find)
    if count != 1:
        return (f"'{find}' was not found in {p}; read_file it first and copy the text exactly." if count == 0
                else f"'{find}' appears {count} times in {p}; include more surrounding text so it's unique.")
    backup(p)
    p.write_text(text.replace(find, replace), encoding="utf-8")
    return f"Edited {p}."


def remember(fact):
    mem["facts"].append(fact)
    save_memory(mem)
    return "Saved to memory."


def _schema(tool_name, desc, /, **params):
    return {"type": "function", "function": {
        "name": tool_name, "description": desc,
        "parameters": {"type": "object", "required": list(params),
                       "properties": {k: {"type": "string", "description": v} for k, v in params.items()}}}}


TOOLS = [
    _schema("open_app", "Open an installed desktop program, e.g. word, excel, notepad, calculator. "
            "Not for websites.", name="App name"),
    _schema("open_url", "Open a website in the browser, e.g. gmail.com, instagram.com.", url="Full URL or domain"),
    _schema("youtube_search", "Search YouTube.", query="Search text"),
    _schema("web_search", "Search Google.", query="Search text"),
    _schema("write_word_doc", "Create a Word document with the given text and open it. "
            "Write the full content yourself. Use '# ' and '## ' for headings, '- ' for bullets.",
            title="Document title", content="Full document text"),
    _schema("draft_email", "Open an email draft (never sends).", to="Recipient address",
            subject="Subject", body="Body text"),
    _schema("list_files", "List what is inside a folder, e.g. Desktop or Documents.", folder="Folder path"),
    _schema("read_file", "Read a text file (txt, md, py, json, csv…).", path="File path"),
    _schema("write_file", "Create a new text file, or replace a whole file's content.",
            path="File path, e.g. Desktop/notes.txt", content="Full file content"),
    _schema("edit_file", "Change part of a text file: replace an exact piece of its text. "
            "Use read_file first so 'find' matches exactly.",
            path="File path", find="Exact existing text", replace="New text"),
    _schema("remember", "Save a fact about the user for later.", fact="The fact"),
    _schema("use_screen", "Look at the screen and click/type to do a task inside an app, "
            "only when no other tool can do it.", task="What to do, step by step"),
    _schema("reply", "Talk to the user: answer a question or chat. NEVER use this to say you opened, "
            "created, saved or did something; call that action's tool instead.", text="What to say"),
]
# Tools whose output the model must read before answering; every other tool's result is reported as-is.
INFO_TOOLS = {"list_files", "read_file"}
TOOL_FUNCS = {"open_app": open_app, "open_url": open_url, "youtube_search": youtube_search,
              "web_search": web_search, "write_word_doc": write_word_doc,
              "draft_email": draft_email, "remember": remember, "list_files": list_files,
              "read_file": read_file, "write_file": write_file, "edit_file": edit_file}


def run_tool(name, raw_args, ui):
    try:
        args = json.loads(raw_args or "{}")
        if name == "use_screen":
            return screen_task(args["task"], ui)
        return TOOL_FUNCS[name](**args)
    except Exception as e:  # report back to the model so it can recover
        return f"Error in {name}: {e}"


# ---------- local agent ----------

def system_prompt():
    sites = "\n".join(f"- {u['when']}: {u['url']}" for u in mem["urls"][-URL_LIMIT:]) or "none"
    facts = "\n".join(f"- {f}" for f in mem["facts"]) or "none"
    return (
        "You are Jarvis, a fast desktop assistant on the user's Windows PC. "
        "Always respond by calling tools. To do something, call its tool; saying it is done does nothing. "
        "Use the reply tool only to talk. Reply briefly in the user's language (Turkish or English). "
        f"Today is {datetime.date.today()}. User folder: {Path.home()} "
        f"(Desktop, Documents, Downloads are inside it). File tools may use: {', '.join(map(str, FILE_ROOTS))}.\n"
        f"Facts about the user:\n{facts}\n"
        f"Sites opened before (newest last), use these for 'the site I mentioned':\n{sites}\n/no_think"
    )


def strip_think(text):
    return re.sub(r"<think>.*?</think>", "", text or "", flags=re.S).strip()


def ask(user_text, ui):
    """ui needs .status(str) and .set_visible(bool).

    The answer is built from what the tools really did, never from the model's own claims:
    small models happily say "Word is open!" without opening anything."""
    turn = [{"role": "user", "content": user_text}]  # saved to history with the real tool calls
    messages = [{"role": "system", "content": system_prompt()}] + recent(mem["history"]) + turn
    said, did, nudged = [], [], False
    for _ in range(MAX_TOOL_ROUNDS):
        msg = local.chat.completions.create(model=LOCAL_MODEL, messages=messages, tools=TOOLS,
                                            tool_choice="required").choices[0].message
        if not msg.tool_calls:  # Ollama ignores tool_choice, so plain text still happens
            text = strip_think(msg.content)
            if not claims_action(text):
                said.append(text)
            elif not nudged:
                nudged = True
                messages.append({"role": "system", "content": "You did not call any tool, so nothing "
                                 "happened. Call the correct tool now."})
                continue
            else:
                said.append(NOTHING_DONE)
            break
        call_msg = msg.model_dump(exclude_none=True)
        messages.append(call_msg)
        turn.append(call_msg)
        needs_followup = False
        for call in msg.tool_calls:
            name = call.function.name
            if name == "reply":
                result = said_text = _reply_text(call.function.arguments)
                said.append(said_text)
            else:
                ui.status(f"Acting: {name}…")
                result = run_tool(name, call.function.arguments, ui)
                if name in INFO_TOOLS and not result.startswith("Error"):
                    needs_followup = True  # the model has to read this and answer
                else:
                    did.append(result)
            messages.append({"role": "tool", "tool_call_id": call.id, "content": result})
            turn.append({"role": "tool", "tool_call_id": call.id, "content": result[:HISTORY_TOOL_CHARS]})
        if said or not needs_followup:
            break
    reply = compose_reply(did, said)
    mem["history"] += turn + [{"role": "assistant", "content": reply}]
    save_memory(mem)
    return reply


def _reply_text(raw_args):
    try:
        return str(json.loads(raw_args or "{}").get("text", "")).strip()
    except json.JSONDecodeError:
        return ""


NOTHING_DONE = "I couldn't do that, nothing was done. Try saying it another way."
# ponytail: keyword check for "I did it" text without a tool call; extend the list if a new phrasing slips through.
CLAIM_RE = re.compile(r"\b(opened|is (now )?open|created|saved|wrote|written|i'?ll remember|i will remember|"
                      r"searched|launched|started)\b|açtım|açıldı|oluşturdum|kaydettim|yazdım|hatırlayacağım|"
                      r"aradım|başlattım", re.I)


def claims_action(text):
    return bool(CLAIM_RE.search(text or ""))


def compose_reply(did, said):
    """Real tool results first, then whatever the model wanted to say."""
    return "\n".join(p for p in [*did, *said] if p) or NOTHING_DONE


# ---------- screen agent (Azure) ----------

SCREEN_PROMPT = """You control a Windows PC by looking at screenshots.
Task: {task}
Screenshot size: {w}x{h} pixels. Reply with ONE JSON object and nothing else:
{{"action":"click","x":0,"y":0}} | {{"action":"double_click","x":0,"y":0}} |
{{"action":"type","text":"..."}} | {{"action":"key","keys":["ctrl","s"]}} |
{{"action":"scroll","amount":-5}} | {{"action":"wait"}} | {{"action":"done","summary":"..."}}
Never press Send on emails, delete files, or buy anything: stop with done and tell the user.
Actions so far: {done}"""


def parse_action(text):
    match = re.search(r"\{.*\}", text or "", flags=re.S)
    if not match:
        raise ValueError(f"No JSON in model reply: {text!r}")
    return json.loads(match.group(0))


def screenshot_b64():
    import pyautogui
    img = pyautogui.screenshot()
    scale = SCREENSHOT_WIDTH / img.width
    img = img.resize((SCREENSHOT_WIDTH, round(img.height * scale)))
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=70)
    return base64.b64encode(buf.getvalue()).decode(), img.width, img.height, scale


def do_action(act, scale):
    import pyautogui
    pyautogui.FAILSAFE = True  # slam the mouse into a screen corner to abort
    kind = act["action"]
    if kind in ("click", "double_click"):
        x, y = round(act["x"] / scale), round(act["y"] / scale)
        (pyautogui.doubleClick if kind == "double_click" else pyautogui.click)(x, y)
    elif kind == "type":
        # pyautogui can't type Turkish letters, so paste through the clipboard.
        import pyperclip
        pyperclip.copy(act["text"])
        pyautogui.hotkey("ctrl", "v")
    elif kind == "key":
        pyautogui.hotkey(*act["keys"])
    elif kind == "scroll":
        pyautogui.scroll(int(act["amount"]) * 100)


def screen_task(task, ui):
    # ponytail: general vision model + JSON clicks is imprecise on dense UIs;
    # swap in a dedicated computer-use model (or local UI-TARS) if clicks keep missing.
    if not azure:
        return "Screen control needs Azure: set AZURE_ENDPOINT, AZURE_API_KEY, AZURE_MODEL in .env."
    import time
    done = []
    ui.set_visible(False)
    try:
        for step in range(MAX_SCREEN_STEPS):
            ui.status(f"Acting on screen… step {step + 1}")
            time.sleep(0.6)
            b64, w, h, scale = screenshot_b64()
            text = SCREEN_PROMPT.format(task=task, w=w, h=h, done=json.dumps(done))
            resp = azure.chat.completions.create(model=AZURE_MODEL, messages=[{"role": "user", "content": [
                {"type": "text", "text": text},
                {"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{b64}"}}]}])
            act = parse_action(resp.choices[0].message.content)
            if act["action"] == "done":
                return act.get("summary", "Done.")
            do_action(act, scale)
            done.append(act)
        return f"Stopped after {MAX_SCREEN_STEPS} steps; task may be unfinished."
    finally:
        ui.set_visible(True)
